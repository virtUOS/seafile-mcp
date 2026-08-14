# /// script
# requires-python = ">=3.11"
# dependencies = [
#   "python-docx>=1.1",
# ]
# ///
"""Generate a genuine Word (.docx) document from a title and plain-text body.

Not part of the deployed package — run standalone with `uv run make_docx.py`,
which installs its own deps (python-docx) without touching anything else.

The body text is split on blank lines into paragraphs. Upload the result with
seafile_upload_file (base64-encode it first, e.g. with file_to_base64.py) —
never with seafile_write_file, which only ever sends UTF-8 text and will
corrupt binary .docx bytes (a .docx is a zip archive, not text).

Example:
    uv run make_docx.py --out /tmp/report.docx --title "Q3 Report" \\
        --text-file /tmp/body.txt
"""

from __future__ import annotations

import argparse


def build_docx(out_path: str, title: str, body: str) -> None:
    import docx

    document = docx.Document()
    document.add_heading(title, level=1)
    for para in body.split("\n\n"):
        para = para.strip()
        if para:
            document.add_paragraph(para)
    document.save(out_path)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", required=True, help="Output .docx path")
    parser.add_argument("--title", required=True, help="Document title (heading)")
    body_group = parser.add_mutually_exclusive_group(required=True)
    body_group.add_argument("--text", help="Body text (paragraphs separated by blank lines)")
    body_group.add_argument("--text-file", help="Path to a file with the body text")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    if args.text_file:
        with open(args.text_file, "r", encoding="utf-8") as f:
            body = f.read()
    else:
        body = args.text

    build_docx(args.out, args.title, body)
    print(f"wrote {args.out}")


if __name__ == "__main__":
    main()
