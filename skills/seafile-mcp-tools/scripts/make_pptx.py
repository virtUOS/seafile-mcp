# /// script
# requires-python = ">=3.11"
# dependencies = [
#   "python-pptx>=1.0",
# ]
# ///
"""Generate a genuine PowerPoint (.pptx) deck from a title slide plus a JSON
list of content slides.

Not part of the deployed package — run standalone with `uv run make_pptx.py`,
which installs its own deps (python-pptx) without touching anything else.

--slides-json must be a JSON array of objects, each shaped like:
    {"title": "Slide heading", "bullets": ["point one", "point two"]}

Upload the result with seafile_upload_file (base64-encode it first, e.g. with
file_to_base64.py) — never with seafile_write_file, which only ever sends
UTF-8 text and will corrupt binary .pptx bytes (a .pptx is a zip archive, not
text).

Example:
    echo '[{"title": "Highlights", "bullets": ["Revenue up 12%", "Churn down"]}]' \\
        > /tmp/slides.json
    uv run make_pptx.py --out /tmp/deck.pptx --title "Q3 Review" \\
        --subtitle "Sales team" --slides-json /tmp/slides.json
"""

from __future__ import annotations

import argparse
import json


def build_pptx(out_path: str, title: str, subtitle: str, slides: list) -> None:
    from pptx import Presentation

    prs = Presentation()

    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = title
    if subtitle and len(title_slide.placeholders) > 1:
        title_slide.placeholders[1].text = subtitle

    for spec in slides:
        heading = spec.get("title", "")
        bullets = [b for b in spec.get("bullets", []) if b]

        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = heading

        if not bullets:
            continue
        body = slide.placeholders[1].text_frame
        body.text = bullets[0]
        for bullet in bullets[1:]:
            body.add_paragraph().text = bullet

    prs.save(out_path)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", required=True, help="Output .pptx path")
    parser.add_argument("--title", required=True, help="Title slide heading")
    parser.add_argument("--subtitle", default="", help="Title slide subtitle")
    parser.add_argument(
        "--slides-json",
        required=True,
        help='Path to a JSON file: [{"title": ..., "bullets": [...]}, ...]',
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    with open(args.slides_json, "r", encoding="utf-8") as f:
        slides = json.load(f)
    if not isinstance(slides, list):
        raise SystemExit("--slides-json must contain a JSON array of slide objects")

    build_pptx(args.out, args.title, args.subtitle, slides)
    print(f"wrote {args.out}")


if __name__ == "__main__":
    main()
