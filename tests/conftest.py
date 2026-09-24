from __future__ import annotations

import asyncio
import os
import zipfile
from io import BytesIO

import pytest
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from seafile_mcp import auth, safety
from seafile_mcp._http import close_http_client
from seafile_mcp.config import get_settings

SERVER = "https://seafile.test"


def pdf_with_pages(count: int, *, texted: bool = False) -> bytes:
    """Build a PDF with `count` pages, either blank or each holding real text.

    Used by both test_documents.py (unit tests on the extraction module) and
    test_server.py (end-to-end tests through the seafile_read_file tool).
    """
    writer = PdfWriter()
    for i in range(count):
        page = writer.add_blank_page(width=200, height=200)
        if texted:
            content = DecodedStreamObject()
            content.set_data(f"BT /F1 12 Tf 10 100 Td (Page {i + 1}) Tj ET".encode())
            content_ref = writer._add_object(content)
            font = DictionaryObject(
                {
                    NameObject("/Type"): NameObject("/Font"),
                    NameObject("/Subtype"): NameObject("/Type1"),
                    NameObject("/BaseFont"): NameObject("/Helvetica"),
                }
            )
            font_ref = writer._add_object(font)
            page[NameObject("/Resources")] = DictionaryObject(
                {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref})}
            )
            page[NameObject("/Contents")] = content_ref
    buf = BytesIO()
    writer.write(buf)
    return buf.getvalue()


def encrypted_pdf(count: int) -> bytes:
    writer = PdfWriter()
    for _ in range(count):
        writer.add_blank_page(width=200, height=200)
    writer.encrypt(user_password="secret")
    buf = BytesIO()
    writer.write(buf)
    return buf.getvalue()


def docx_with_paragraphs(paragraphs: list[str], *, tables: list[list[list[str]]] = ()) -> bytes:
    """Build a minimal .docx with the given paragraph texts and table rows."""
    from docx import Document

    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    for rows in tables:
        table = doc.add_table(rows=len(rows), cols=len(rows[0]))
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def pptx_with_slides(slide_texts: list[str]) -> bytes:
    """Build a minimal .pptx with one title-only slide per entry in `slide_texts`."""
    from pptx import Presentation

    prs = Presentation()
    layout = prs.slide_layouts[5]  # "Title Only"
    for text in slide_texts:
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = text
    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


def xlsx_with_sheets(sheets: dict[str, list[list[object]]]) -> bytes:
    """Build a minimal .xlsx with one sheet per (name, rows) entry in `sheets`."""
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(title=name)
        for row in rows:
            ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def image_bytes(
    width: int = 40,
    height: int = 30,
    fmt: str = "PNG",
    *,
    mode: str = "RGB",
    noise: bool = False,
) -> bytes:
    """Build a real image of the given size and format.

    Used by test_documents.py (unit tests on is_image/render_image) and by
    test_server.py and test_client.py end-to-end. `noise` fills the image with
    random pixels, which defeats every compressor — the only way to build a
    source that actually exercises the byte-cap ladder at a small size.
    """
    from PIL import Image as PILImage

    if noise:
        img = PILImage.frombytes(
            "RGB", (width, height), os.urandom(width * height * 3)
        )
    else:
        img = PILImage.new("RGB", (width, height), (10, 120, 200))
    # Built in RGB and converted, rather than filled directly: PILImage.new
    # wants a mode-appropriate colour, and "a blue square" is not expressible
    # as one for L, 1, CMYK and P all at once.
    if mode == "RGBA":
        img = img.convert("RGBA")
        img.putalpha(90)
    elif mode != "RGB":
        img = img.convert(mode)
    buf = BytesIO()
    img.save(buf, fmt)
    return buf.getvalue()


def animated_gif_bytes(frames: int = 3) -> bytes:
    """Build an animated GIF with `frames` visibly distinct frames.

    They have to differ: Pillow collapses identical frames on save, and a GIF
    that reports n_frames == 1 tests nothing about the first-frame rule.
    """
    from PIL import Image as PILImage

    colors = [(255, 0, 0), (0, 255, 0), (0, 0, 255), (255, 255, 0), (0, 255, 255)]
    images = [
        PILImage.new("RGB", (20, 20), colors[i % len(colors)]) for i in range(frames)
    ]
    buf = BytesIO()
    images[0].save(buf, "GIF", save_all=True, append_images=images[1:])
    return buf.getvalue()


def exif_rotated_jpeg(width: int, height: int, orientation: int = 6) -> bytes:
    """Build a JPEG whose EXIF says it is rotated.

    Orientation 6 means "rotate 90° clockwise to display", so a correctly
    handled render comes back with the axes swapped relative to the stored
    raster. This is what every phone photo looks like.
    """
    from PIL import Image as PILImage

    img = PILImage.new("RGB", (width, height), (10, 120, 200))
    exif = img.getexif()
    exif[0x0112] = orientation
    buf = BytesIO()
    img.save(buf, "JPEG", exif=exif)
    return buf.getvalue()


def oversized_png_header(width: int, height: int) -> bytes:
    """A valid PNG header declaring a huge size, with no pixel data.

    Lets a test drive the decompression-bomb guard without building the
    gigabytes such an image would really occupy — which is the whole point of
    checking the declared size before decoding.
    """
    import struct
    import zlib

    def chunk(kind: bytes, payload: bytes) -> bytes:
        body = kind + payload
        return (
            struct.pack(">I", len(payload))
            + body
            + struct.pack(">I", zlib.crc32(body))
        )

    ihdr = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr) + chunk(b"IEND", b"")


#: A line of German that exercises every byte a cp1252 export gets wrong: the
#: three umlauts, the sharp s, and the Euro sign — the character that tells
#: cp1252 (0x80) apart from true Latin-1, which has no Euro at all.
GERMAN_TEXT = "Prüfungsordnung für Größe\nStraße 5\tGebühr: 12,50 €\n"

#: ODF media types. Repeated here rather than imported from documents.py: a
#: test that builds its input from the same constant the code matches on
#: cannot catch that constant being wrong.
ODT_MIME = "application/vnd.oasis.opendocument.text"
ODS_MIME = "application/vnd.oasis.opendocument.spreadsheet"
ODP_MIME = "application/vnd.oasis.opendocument.presentation"

#: The namespace declarations every ODF body fragment below needs.
ODF_NS = (
    'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
    'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
    'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
    'xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0" '
    'xmlns:presentation="urn:oasis:names:tc:opendocument:xmlns:presentation:1.0"'
)


def odf_bytes(
    media_type: str,
    body_xml: str,
    *,
    manifest: str | None = None,
    mimetype_first: bool = True,
    include_mimetype: bool = True,
    content_override: bytes | None = None,
) -> bytes:
    """Build a minimal ODF package around the given office:body XML.

    Hand-built rather than written with an ODF library, because this project
    has none — the same reason oversized_png_header exists. The three parts
    below are all documents.py ever consults, and the real LibreOffice files in
    docs_format_test_/ confirm the shape. The keyword arguments exist so a test
    can build the packages LibreOffice does *not* write: re-zipped so mimetype
    is no longer first, missing it entirely, password-protected, or corrupt.
    """
    content = content_override if content_override is not None else (
        f'<?xml version="1.0"?><office:document-content {ODF_NS}>'
        f"{body_xml}</office:document-content>"
    ).encode()

    buf = BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        if include_mimetype and mimetype_first:
            info = zipfile.ZipInfo("mimetype")
            info.compress_type = zipfile.ZIP_STORED
            zf.writestr(info, media_type)
        zf.writestr(
            "META-INF/manifest.xml",
            manifest or '<?xml version="1.0"?><manifest:manifest '
            'xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0"/>',
        )
        zf.writestr("content.xml", content)
        if include_mimetype and not mimetype_first:
            zf.writestr("mimetype", media_type)
    return buf.getvalue()


def odt_bytes(body_inner: str, **kw) -> bytes:
    """An .odt whose office:text holds `body_inner`."""
    return odf_bytes(
        ODT_MIME, f"<office:body><office:text>{body_inner}</office:text></office:body>", **kw
    )


def ods_bytes(sheets: dict[str, list[list[str]]] | None = None, raw: str = "", **kw) -> bytes:
    """An .ods built from {sheet name: rows of cell text}, or from raw XML.

    `raw` is the escape hatch for the repeat-count and padding cases, which are
    about attributes no cell-text mapping can express.
    """
    if raw:
        body = raw
    else:
        body = ""
        for name, rows in (sheets or {}).items():
            cells = "".join(
                "<table:table-row>"
                + "".join(
                    f"<table:table-cell><text:p>{c}</text:p></table:table-cell>"
                    for c in row
                )
                + "</table:table-row>"
                for row in rows
            )
            body += f'<table:table table:name="{name}">{cells}</table:table>'
    return odf_bytes(
        ODS_MIME,
        f"<office:body><office:spreadsheet>{body}</office:spreadsheet></office:body>",
        **kw,
    )


def odp_bytes(slide_texts: list[str], *, notes: str | None = None, **kw) -> bytes:
    """An .odp with one draw:page per entry, optionally with speaker notes.

    A draw:master-page is always included: a real presentation has one as a
    sibling of the slides (confirmed against a LibreOffice file), and counting
    it would invent a slide nobody wrote.
    """
    pages = '<draw:master-page draw:name="Default"/>'
    for i, text in enumerate(slide_texts, start=1):
        note_xml = (
            f"<presentation:notes><draw:frame><draw:text-box><text:p>{notes}"
            f"</text:p></draw:text-box></draw:frame></presentation:notes>"
            if notes
            else ""
        )
        pages += (
            f'<draw:page draw:name="Slide {i}">'
            f"<draw:frame><draw:text-box><text:p>{text}</text:p></draw:text-box>"
            f"</draw:frame>{note_xml}</draw:page>"
        )
    return odf_bytes(
        ODP_MIME,
        f"<office:body><office:presentation>{pages}</office:presentation></office:body>",
        **kw,
    )


#: A manifest declaring the package encrypted, as LibreOffice writes when a
#: password is set. The manifest itself stays in the clear because it carries
#: the key-derivation parameters.
ENCRYPTED_MANIFEST = (
    '<?xml version="1.0"?><manifest:manifest '
    'xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0">'
    '<manifest:file-entry manifest:full-path="content.xml">'
    "<manifest:encryption-data/></manifest:file-entry></manifest:manifest>"
)


#: Minimal bytes recognized as an OLE2/CFB container (legacy or encrypted
#: Office file) by documents.is_legacy_or_encrypted_office. The magic number
#: alone is enough; no real encrypted file needed.
CFB_MAGIC_BYTES = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 32


@pytest.fixture(autouse=True)
def _env(monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setenv("SEAFILE_SERVER_URL", SERVER)
    monkeypatch.delenv("SEAFILE_API_TOKEN", raising=False)
    monkeypatch.delenv("SEAFILE_MCP_MODE", raising=False)
    monkeypatch.delenv("SEAFILE_MCP_ALLOWED_EMAIL_DOMAINS", raising=False)
    get_settings.cache_clear()
    auth.clear_mode_cache()
    safety.reset_state()
    yield
    asyncio.run(close_http_client())
    get_settings.cache_clear()
    auth.clear_mode_cache()
    safety.reset_state()


@pytest.fixture
def no_http_headers(monkeypatch: pytest.MonkeyPatch):
    """Simulate stdio / no active HTTP request."""
    monkeypatch.setattr(auth, "get_http_headers", lambda **kw: {})


@pytest.fixture
def headers(monkeypatch: pytest.MonkeyPatch):
    """Let a test set the headers seen by the auth layer."""
    current: dict[str, str] = {}

    def _set(**kw: str) -> None:
        current.clear()
        current.update({k.replace("_", "-").lower(): v for k, v in kw.items()})

    monkeypatch.setattr(auth, "get_http_headers", lambda **kw: dict(current))
    return _set
